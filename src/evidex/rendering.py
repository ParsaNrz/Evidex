"""Application-owned document and audit rendering; never execute model markup."""

from __future__ import annotations

import hashlib
import json
import os
import re
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from fpdf import FPDF

from evidex.models import (
    CitationAuditEntry,
    CitationStyle,
    ExportedFile,
    ExportPackage,
    Language,
    ProjectBrief,
    SectionDraft,
    SourceRegistry,
)

DISCLOSURE = (
    "AI-assisted draft. A language model assisted with planning, organization, and prose. "
    "The student must verify claims, source relevance, references, and compliance with institutional "
    "AI policies. DOI metadata verification does not establish claim entailment. "
    "No original research data or findings have been generated."
)
DEMO_WARNING = (
    "SYNTHETIC DEMO — fixture sources, not evidence. Do not submit or cite this document."
)


class DocumentRenderer:
    def __init__(self, root: str | Path, font_path: str | Path | None = None):
        self.root = Path(root).resolve()
        self.font_path = Path(font_path or "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")

    def render(self, job_id: str, payload: dict, *, demo_mode: bool = False) -> ExportPackage:
        brief = ProjectBrief.model_validate(payload["brief"])
        registry = SourceRegistry.model_validate(payload["registry"])
        sections = [SectionDraft.model_validate(s) for s in payload["sections"]]
        sources = {s.source_id: s for s in registry.sources}
        claims: dict[str, list[str]] = {}
        notes: dict[str, set[str]] = {}
        for section in sections:
            for claim in section.claims:
                for sid in claim.source_ids:
                    if sid not in sources:
                        raise ValueError("Cannot render an unregistered citation")
                    claims.setdefault(sid, []).append(claim.claim_id)
                    notes.setdefault(sid, set()).update(
                        n
                        for n in claim.evidence_note_ids
                        if n in {n.note_id for n in sources[sid].notes}
                    )
        ordered = list(claims)
        labels = {}
        references = []
        for index, sid in enumerate(ordered, 1):
            source = sources[sid]
            if brief.citation_style == CitationStyle.APA7:
                # Include ID for unambiguous disambiguation; bibliography is a metadata draft.
                labels[sid] = f"({source.authors[0]}, {source.year}; {sid})"
            else:
                labels[sid] = f"[{index}]"
            references.append(
                f"[{index}] {', '.join(source.authors)} ({source.year}). {source.title}. "
                + (f"{source.venue}. " if source.venue else "")
                + (source.url or source.doi or "")
                + f" [Registry: {sid}]"
            )

        def citations(text: str) -> str:
            def replace(match):
                if match[1] not in labels:
                    raise ValueError("Inline citation has no audited claim")
                return labels[match[1]]

            return re.sub(r"\[@([^\]]+)\]", replace, text)

        disclosure = (DEMO_WARNING + "\n\n" if demo_mode else "") + DISCLOSURE
        paragraphs: list[tuple[str, str]] = [("title", brief.topic), ("text", disclosure)]
        for section in sections:
            paragraphs.append(("heading", section.title))
            paragraphs.extend(
                ("text", citations(p)) for p in section.body.splitlines() if p.strip()
            )
        paragraphs.append(("heading", "منابع / References"))
        paragraphs.append(
            (
                "text",
                "References are a metadata-based draft; verify formatting against "
                f"{brief.citation_style.value} and your university requirements.",
            )
        )
        paragraphs.extend(("text", ref) for ref in references)
        paragraphs.append(("heading", "Citation audit / ممیزی استناد"))
        for sid in ordered:
            paragraphs.append(
                (
                    "text",
                    f"{sid}: {sources[sid].verification.status.value}; "
                    f"claims: {', '.join(claims[sid])}; notes: {', '.join(sorted(notes[sid]))}",
                )
            )
        if payload.get("overrides"):
            paragraphs.append(("heading", "Human overrides"))
            paragraphs.extend(
                ("text", json.dumps(o, ensure_ascii=False)) for o in payload["overrides"]
            )
        package = ExportPackage(
            citation_audit=[
                CitationAuditEntry(
                    source_id=sid,
                    verification_status=sources[sid].verification.status,
                    claim_ids=claims[sid],
                    evidence_note_ids=sorted(notes[sid]),
                )
                for sid in ordered
            ],
            change_log=[
                "Rendered approved, reviewed sections with registered citations.",
                "Reference formatting is a metadata draft, not certified style compliance.",
            ],
            human_overrides=payload.get("overrides", []),
            ai_use_disclosure=disclosure,
            audit_complete=True,
            demo_mode=demo_mode,
        )
        # Model/user text never determines filesystem paths. Replace complete files atomically.
        directory = self.root / hashlib.sha256(job_id.encode()).hexdigest()
        directory.mkdir(parents=True, exist_ok=True)
        if not self.font_path.is_file():
            raise ValueError("PDF font missing; set EVIDEX_PDF_FONT to a Unicode TTF font")
        with TemporaryDirectory(dir=directory) as temp:
            temp = Path(temp)
            self._docx(paragraphs, temp / "document.docx", brief.language)
            self._pdf(paragraphs, temp / "document.pdf", brief.language)
            (temp / "audit.json").write_text(
                json.dumps(
                    {
                        "package": package.model_dump(mode="json"),
                        "registry": registry.model_dump(mode="json"),
                        "sections": [s.model_dump(mode="json") for s in sections],
                        "review": payload.get("review"),
                    },
                    ensure_ascii=False,
                    indent=2,
                ),
                encoding="utf-8",
            )
            for name, fmt in [
                ("document.docx", "docx"),
                ("document.pdf", "pdf"),
                ("audit.json", "json"),
            ]:
                destination = directory / name
                os.replace(temp / name, destination)
                package.files.append(
                    ExportedFile(
                        format=fmt, path=str(destination), size_bytes=destination.stat().st_size
                    )
                )
        return package

    @staticmethod
    def _docx(paragraphs, path: Path, language: Language):
        doc = Document()
        normal = doc.styles["Normal"]
        normal.font.name = "DejaVu Sans"
        normal.font.size = Pt(11)
        for kind, text in paragraphs:
            p = (
                doc.add_heading(text, 0 if kind == "title" else 1)
                if kind != "text"
                else doc.add_paragraph(text)
            )
            if language == Language.PERSIAN:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                bidi = OxmlElement("w:bidi")
                p._p.get_or_add_pPr().append(bidi)
            for run in p.runs:
                run.font.name = "DejaVu Sans"
                props = run._r.get_or_add_rPr()
                fonts = props.rFonts
                if fonts is not None:
                    fonts.set(qn("w:cs"), "DejaVu Sans")
                if language == Language.PERSIAN:
                    rtl = OxmlElement("w:rtl")
                    props.append(rtl)
        doc.save(path)

    def _pdf(self, paragraphs, path: Path, language: Language):
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=18)
        pdf.set_margins(18, 18, 18)
        pdf.add_font("Unicode", fname=str(self.font_path))
        pdf.set_text_shaping(True, direction="rtl" if language == Language.PERSIAN else "ltr")
        pdf.add_page()
        for kind, text in paragraphs:
            size = {"title": 19, "heading": 14, "text": 10}[kind]
            pdf.set_font("Unicode", size=size)
            pdf.multi_cell(
                0,
                size * 0.55,
                text,
                new_x="LMARGIN",
                new_y="NEXT",
                align="R" if language == Language.PERSIAN else "L",
            )
            pdf.ln(3)
        pdf.output(path)
